import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2].parent
SOURCE = ROOT / "docs" / "MELHORIAS.md"
DEFAULT_OUTPUT = ROOT / "docs" / "MELHORIAS.docx"


def set_font(run, size=12, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_paragraph(document, text, size=12, bold=False, italic=False, before=0, after=3):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)
    return paragraph


def is_section_heading(line):
    if not line or line.startswith("- [") or line[0].isdigit():
        return False
    return line == line.upper() and any(character.isalpha() for character in line)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(12)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    for index, raw_line in enumerate(lines):
        line = raw_line.strip()
        if not line:
            continue

        if index == 0:
            paragraph = add_paragraph(document, line, size=14, bold=True, after=10)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue

        if line.startswith("Formatação da versão"):
            add_paragraph(document, line, size=12, italic=True, after=8)
            continue

        if line.startswith("==========================================="):
            add_paragraph(document, "MÉDIO PRAZO", size=14, bold=True, before=8, after=6)
            continue

        if is_section_heading(line):
            add_paragraph(document, line, size=14, bold=True, before=10, after=6)
            continue

        add_paragraph(document, line, size=12, after=3)

    document.core_properties.title = "Drowned - Melhorias e Próximos Passos"
    document.core_properties.subject = "Backlog de melhorias do protótipo e do prólogo"
    document.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
